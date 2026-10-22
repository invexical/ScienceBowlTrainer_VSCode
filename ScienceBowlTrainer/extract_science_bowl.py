#!/usr/bin/env python3
"""Extract Science Bowl questions from locally downloaded Drive folder documents.

Dependencies: PyMuPDF (pymupdf), python-docx. Works with PDFs, DOCX, TXT, MD.
No internet connection or API keys required.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document

SUPPORTED = {'.pdf', '.docx', '.txt', '.md'}
IGNORED = {'.ds_store', 'thumbs.db'}

CATEGORY_NAMES = [
    ('Earth & Space Science', r'(?:EARTH\s*(?:AND|&)\s*SPACE(?:\s*SCIENCE)?|EARTH\s*SCIENCE|ASTRONOMY)'),
    ('Computer Science', r'(?:COMPUTER\s*SCIENCE|COMPUTING)'),
    ('General Science', r'(?:GENERAL\s*SCIENCE)'),
    ('Biology', r'(?:BIOLOGY|LIFE\s*SCIENCE)'),
    ('Chemistry', r'(?:CHEMISTRY)'),
    ('Mathematics', r'(?:MATHEMATICS|MATH)'),
    ('Physics', r'(?:PHYSICS)'),
    ('Energy', r'(?:ENERGY)'),
]
CATEGORIES = [(name, re.compile(r'^\s*' + pat + r'\b\s*', re.I)) for name, pat in CATEGORY_NAMES]
SECTION = re.compile(r'^\s*(?:(VISUAL)\s+)?(TOSS[\s-]*UP|BONUS)(?:\s+(VISUAL))?\s*[:\-–—]?\s*(.*)$', re.I)
NUMBER = re.compile(r'^\s*(\d{1,3})\s*[).:]\s*(.*)$')
FORMAT = re.compile(r'^\s*(MULTIPLE[\s-]*CHOICE|SHORT[\s-]*ANSWER|MCQ|MC|SA)\b\s*', re.I)
ANSWER = re.compile(r'^\s*(?:CORRECT\s+)?ANSWERS?\s*[:\-–—]\s*(.*)$', re.I)
CHOICE = re.compile(r'^\s*([WXYZ])\s*[).:]\s*(.+)$', re.I)
ROUND = re.compile(r'\bround[\s_-]*(\d{1,3})\b', re.I)
YEAR = re.compile(r'\b(20\d{2})\b')
PAGE_NO = re.compile(r'^\s*(?:page\s*\d+|\d+\s*/\s*\d+)\s*$', re.I)
ROUND_PAGE = re.compile(r'^.*\bscience\s*bowl.*\bround\s*\d+\s+page\s*\d+\s*$', re.I)
ONLY_TITLE = re.compile(r'^\s*(?:20\d{2}\s+)?(?:MIT\s+)?Science\s+Bowl\s+(?:High\s+School\s+)?Invitational\s*$', re.I)

FIELDS = [
    'id', 'year', 'tournament', 'round', 'level', 'number', 'type', 'category',
    'format', 'question', 'choice_w', 'choice_x', 'choice_y', 'choice_z',
    'answer', 'source_file', 'source_page', 'has_image_on_page',
    'needs_review', 'review_reasons',
]


def normalize(value: str) -> str:
    return re.sub(r'\s+', ' ', value).strip()


def looks_like_header(text: str) -> bool:
    return any(pattern.match(text) for _, pattern in CATEGORIES)


def parse_lead(text: str) -> tuple[str, str, str]:
    """Extract category and answer format at start of a question's first line."""
    text = text.lstrip(' \t:–—-?')
    category = ''
    for name, pattern in CATEGORIES:
        m = pattern.match(text)
        if m:
            category = name
            text = text[m.end():].lstrip(' \t:–—-?')
            break
    m = FORMAT.match(text)
    fmt = ''
    if m:
        fmt = 'Multiple Choice' if m.group(1).upper().replace('-', '').replace(' ', '') in ('MULTIPLECHOICE', 'MCQ', 'MC') else 'Short Answer'
        text = text[m.end():].lstrip(' \t:–—-?')
    return category, fmt, text


def skip_line(line: str) -> bool:
    return bool(PAGE_NO.match(line) or ROUND_PAGE.match(line) or ONLY_TITLE.match(line))


@dataclass
class Pending:
    number: int
    kind: str
    category: str
    format: str
    question_parts: list[str]
    answer_parts: list[str]
    choices: dict[str, str]
    page: int | None
    visual: bool = False
    collecting_answer: bool = False
    last_choice: str = ''


def document_lines(path: Path) -> tuple[list[tuple[str, int | None]], set[int], str]:
    """Return (text lines and source-page numbers, PDF pages with images, raw text)."""
    lines: list[tuple[str, int | None]] = []
    pages_with_images: set[int] = set()
    if path.suffix.lower() == '.pdf':
        doc = pymupdf.open(path)
        try:
            for page_index, page in enumerate(doc):
                page_num = page_index + 1
                if page.get_images(full=True):
                    pages_with_images.add(page_num)
                # PyMuPDF's sorted text works for most single-column Science Bowl PDFs.
                for line in page.get_text('text', sort=True).splitlines():
                    lines.append((line, page_num))
        finally:
            doc.close()
    elif path.suffix.lower() == '.docx':
        doc = Document(path)
        if hasattr(doc, 'iter_inner_content'):
            from docx.table import Table
            for item in doc.iter_inner_content():
                if isinstance(item, Table):
                    for row in item.rows:
                        for cell in row.cells:
                            for line in cell.text.splitlines():
                                lines.append((line, None))
                else:
                    lines.extend((line, None) for line in item.text.splitlines())
        else:
            for p in doc.paragraphs:
                lines.extend((line, None) for line in p.text.splitlines())
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        lines.extend((line, None) for line in cell.text.splitlines())
    else:
        text = path.read_text(encoding='utf-8-sig', errors='replace')
        lines = [(line, None) for line in text.splitlines()]
    raw = '\n'.join(s for s, _ in lines)
    return lines, pages_with_images, raw


def parse_metadata(relative_path: Path, raw: str) -> tuple[str, str, str, str]:
    pathstr = str(relative_path).replace('\\', '/')
    y = YEAR.search(pathstr) or YEAR.search(raw[:2500])
    r = ROUND.search(pathstr) or ROUND.search(raw[:1500])
    # The folder above a round file is generally the most useful tournament label.
    tournament = relative_path.parent.name if relative_path.parent != Path('.') else relative_path.stem
    level = ''
    if re.search(r'\b(middle[ -]?school|\bMS\b)\b', pathstr, re.I):
        level = 'Middle School'
    elif re.search(r'\b(high[ -]?school|\bHS\b)\b', pathstr, re.I):
        level = 'High School'
    return (y.group(1) if y else '', tournament, r.group(1) if r else '', level)


def parse_questions(lines: list[tuple[str, int | None]], image_pages: set[int], relative: Path, raw: str) -> list[dict]:
    year, tournament, round_number, level = parse_metadata(relative, raw)
    records: list[dict] = []
    current: Pending | None = None
    active_section = ''
    active_visual = False
    ready_after_section = False

    def finish() -> None:
        nonlocal current
        if current is None:
            return
        question = normalize(' '.join(current.question_parts))
        answer = normalize(' '.join(current.answer_parts))
        choices = {key: normalize(val) for key, val in current.choices.items()}
        reasons = []
        if not question:
            reasons.append('empty question')
        if not answer:
            reasons.append('missing answer (possibly separate answer key)')
        if not current.category:
            reasons.append('unrecognized category')
        fmt = current.format
        if not fmt and choices:
            fmt = 'Multiple Choice'
        if not fmt:
            reasons.append('unrecognized question format')
        if fmt == 'Multiple Choice' and len(choices) != 4:
            reasons.append('expected four W/X/Y/Z choices')
        if current.visual:
            reasons.append('visual question; inspect associated diagram')
        if current.page in image_pages:
            reasons.append('source PDF page contains an image; check whether needed')
        if len(question) < 15:
            reasons.append('unusually short question; check parsing')
        if len(answer) > 600:
            reasons.append('unusually long answer; check parsing boundary')
        rowid = hashlib.sha1(f'{relative.as_posix()}|{current.number}|{current.kind}|{len(records)}'.encode('utf-8')).hexdigest()[:16]
        records.append({
            'id': rowid,
            'year': year,
            'tournament': tournament,
            'round': round_number,
            'level': level,
            'number': current.number,
            'type': current.kind,
            'category': current.category or 'Unknown',
            'format': fmt or 'Unknown',
            'question': question,
            'choice_w': choices.get('W', ''),
            'choice_x': choices.get('X', ''),
            'choice_y': choices.get('Y', ''),
            'choice_z': choices.get('Z', ''),
            'answer': answer,
            'source_file': relative.as_posix(),
            'source_page': current.page if current.page is not None else '',
            'has_image_on_page': bool(current.page in image_pages),
            'needs_review': bool(reasons),
            'review_reasons': '; '.join(reasons),
        })
        current = None

    for original, page_num in lines:
        line = normalize(original)
        if not line or skip_line(line):
            continue

        sec = SECTION.match(line)
        if sec:
            finish()
            active_section = 'Toss-up' if 'TOSS' in sec.group(2).upper() else 'Bonus'
            active_visual = bool(sec.group(1) or sec.group(3))
            ready_after_section = True
            line = sec.group(4).strip()
            if not line:
                continue

        numbered = NUMBER.match(line)
        if numbered and active_section:
            number, rest = int(numbered.group(1)), numbered.group(2)
            # If this line is a choice-free list (e.g. "1) Pillow lavas"), it
            # must NOT create a question in the middle of an unanswered bonus.
            can_start = (ready_after_section or current is None or
                         bool(current.answer_parts) or looks_like_header(rest))
            if can_start:
                finish()
                category, fmt, lead = parse_lead(rest)
                current = Pending(number, active_section, category, fmt,
                                  [lead] if lead else [], [], {}, page_num, active_visual)
                ready_after_section = False
                continue

        if current is None:
            continue
        ans = ANSWER.match(line)
        if ans:
            current.collecting_answer = True
            if ans.group(1):
                current.answer_parts.append(ans.group(1))
            continue
        if current.collecting_answer:
            current.answer_parts.append(line)
            continue
        option = CHOICE.match(line)
        if option and (current.format == 'Multiple Choice' or not current.format or current.choices):
            key = option.group(1).upper()
            # Duplicate option labels within a question are unusual; keep both texts.
            if key in current.choices:
                current.choices[key] += ' ' + option.group(2)
            else:
                current.choices[key] = option.group(2)
            current.last_choice = key
            continue
        if current.last_choice and current.format == 'Multiple Choice':
            # Wrapped choice text belongs to its most recent option, not the question.
            current.choices[current.last_choice] += ' ' + line
            continue
        # Some PDFs put the category or format on the next line.
        if current.category and not current.format and not current.question_parts:
            f = FORMAT.match(line.lstrip(' \t:–—-?'))
            if f:
                category, fmt, lead = parse_lead(line)
                if not fmt:
                    fmt = 'Multiple Choice' if 'CHOICE' in f.group(1).upper() or f.group(1).upper() in ('MC', 'MCQ') else 'Short Answer'
                    lead = line.lstrip(' \t:–—-?')[f.end():].strip()
                current.format = fmt
                if lead:
                    current.question_parts.append(lead)
                continue
        # Some PDFs have category and format on a second line after the number.
        if not current.question_parts and not current.category:
            category, fmt, lead = parse_lead(line)
            if category or fmt:
                current.category = category
                current.format = fmt
                if lead:
                    current.question_parts.append(lead)
                continue
        current.question_parts.append(line)

    finish()
    return records


def save_csv(path: Path, records: list[dict], columns: list[str]) -> None:
    with path.open('w', encoding='utf-8-sig', newline='') as fp:
        writer = csv.DictWriter(fp, fieldnames=columns, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(records)


def main() -> int:
    p = argparse.ArgumentParser(description='Extract Science Bowl questions from PDFs/DOCX recursively')
    p.add_argument('--input', default='input', type=Path, help='Downloaded and extracted folder')
    p.add_argument('--output', default='output', type=Path, help='Where to save CSV and JSON')
    args = p.parse_args()
    root = args.input.expanduser().resolve()
    out = args.output.expanduser().resolve()
    if not root.is_dir():
        print(f'ERROR: Input folder does not exist: {root}', file=sys.stderr)
        return 2
    if out == root or root in out.parents:
        print('ERROR: Output should not be inside the input folder (avoid rescanning output).', file=sys.stderr)
        return 2
    out.mkdir(parents=True, exist_ok=True)
    all_records: list[dict] = []
    files_report: list[dict] = []
    seen_supported = 0
    for file in sorted(root.rglob('*')):
        if not file.is_file() or file.name.lower() in IGNORED:
            continue
        relative = file.relative_to(root)
        suffix = file.suffix.lower()
        if suffix not in SUPPORTED:
            files_report.append({'source_file': relative.as_posix(), 'status': 'unsupported format', 'questions': 0, 'details': suffix or '(no extension)'})
            continue
        seen_supported += 1
        print(f'Processing: {relative}')
        try:
            lines, image_pages, raw = document_lines(file)
            if not raw.strip():
                status = 'no extractable text (scanned PDF or empty file)'
                files_report.append({'source_file': relative.as_posix(), 'status': status, 'questions': 0, 'details': ''})
                continue
            records = parse_questions(lines, image_pages, relative, raw)
            all_records.extend(records)
            status = 'parsed' if records else 'no recognizable questions'
            notes = []
            for kind in ('Toss-up', 'Bonus'):
                nums = [r['number'] for r in records if r['type'] == kind]
                if not nums:
                    if records:
                        notes.append(f'no {kind.lower()} questions detected')
                    continue
                gaps = sorted(set(range(1, max(nums) + 1)) - set(nums))
                if gaps:
                    notes.append(f'{kind.lower()} numbering gaps: {gaps}')
                repeated = sorted(n for n, count in Counter(nums).items() if count > 1)
                if repeated:
                    notes.append(f'duplicate {kind.lower()} numbers: {repeated}')
            files_report.append({'source_file': relative.as_posix(), 'status': status, 'questions': len(records), 'details': '; '.join(notes)})
        except Exception as exc:
            files_report.append({'source_file': relative.as_posix(), 'status': 'error', 'questions': 0, 'details': f'{type(exc).__name__}: {exc}'})
    # Preserve every occurrence: repeated round PDFs are not removed silently.
    duplicates = []
    first_seen: dict[str, str] = {}
    for record in all_records:
        fingerprint = hashlib.sha256((normalize(record['question']).casefold() + '\\' + normalize(record['answer']).casefold()).encode('utf-8')).hexdigest()
        if fingerprint in first_seen:
            duplicates.append({'duplicate_id': record['id'], 'original_id': first_seen[fingerprint], 'source_file': record['source_file']})
        else:
            first_seen[fingerprint] = record['id']
    save_csv(out / 'questions.csv', all_records, FIELDS)
    (out / 'questions.json').write_text(json.dumps(all_records, ensure_ascii=False, indent=2), encoding='utf-8')
    issues = [r for r in all_records if r['needs_review']]
    save_csv(out / 'questions_needing_review.csv', issues, FIELDS)
    save_csv(out / 'file_report.csv', files_report, ['source_file', 'status', 'questions', 'details'])
    save_csv(out / 'possible_duplicates.csv', duplicates, ['duplicate_id', 'original_id', 'source_file'])
    summary = {
        'total_files_seen': len(files_report),
        'supported_files': seen_supported,
        'files_by_status': dict(Counter(r['status'] for r in files_report)),
        'questions_extracted': len(all_records),
        'questions_needing_review': len(issues),
        'possible_duplicate_questions': len(duplicates),
        'by_category': dict(Counter(r['category'] for r in all_records)),
        'by_type': dict(Counter(r['type'] for r in all_records)),
        'by_format': dict(Counter(r['format'] for r in all_records)),
        'output_files': ['questions.csv', 'questions.json', 'questions_needing_review.csv',
                         'file_report.csv', 'possible_duplicates.csv', 'summary.json'],
    }
    (out / 'summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print('\n=== RESULTS ===')
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f'\nSaved results to {out}')
    if issues or any(r['status'] != 'parsed' for r in files_report):
        print('IMPORTANT: Review questions_needing_review.csv and file_report.csv for missed/malformed documents.')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
